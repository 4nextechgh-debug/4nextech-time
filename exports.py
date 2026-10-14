"""Excel and PDF exports of the attendance report, branded 4nextech Time.

openpyxl and reportlab are imported inside the functions so the rest of the
app still runs if either library is missing.
"""
import io
import os
from datetime import datetime

BRAND = "4nextech"
PRODUCT = "Time"
BLUE = "1F3A8A"
GREY = "667085"
RED = "C0392B"
FLAG = ("Absent", "Late", "Missing out")

SUMMARY_COLS = [
    ("Code", "code"), ("Name", "name"), ("Department", "department"),
    ("Days present", "present"), ("Late days", "late"), ("Absent days", "absent"),
    ("Total time (h:mm)", "total_time"), ("Hours worked", "hours"), ("Overtime (h:mm)", "ot_time"),
]
DETAIL_COLS = [
    ("Date", "date"), ("Code", "code"), ("Name", "name"), ("Department", "department"),
    ("Shift", "shift"), ("Clock in", "first_in"), ("Clock out", "last_out"), ("Total time", "total_time"), ("Hours", "worked_hours"),
    ("Late (min)", "late_min"), ("Overtime (min)", "overtime_min"), ("Status", "status"),
]


def _period(start, end, dept):
    text = f"{start:%d %b %Y} to {end:%d %b %Y}"
    return text + (f"   |   Department: {dept}" if dept else "")


def _summary_rows(summary):
    """Per-user rows plus a grand-total row at the bottom."""
    out = [dict(s) for s in summary]
    if out:
        from reports import hm
        out.append({"code": "", "name": "TOTAL (all users)", "department": "",
                    "present": sum(s["present"] for s in summary), "late": sum(s["late"] for s in summary),
                    "absent": sum(s["absent"] for s in summary),
                    "total_time": hm(sum(s["minutes"] for s in summary)),
                    "hours": round(sum(s["minutes"] for s in summary) / 60, 2),
                    "ot_time": hm(sum(s["overtime"] for s in summary))})
    return out


# ---------------------------------------------------------------- Excel
def xlsx_bytes(rows, summary, start, end, dept=None, company=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor=BLUE)
    head_font = Font(bold=True, color="FFFFFF")
    header_row = 5

    def fill_sheet(ws, title, cols, data):
        ws.title = title
        ws["A1"] = company or BRAND
        ws["A1"].font = Font(bold=True, size=22, color=BLUE)
        ws["A1"].data_type = "s"  # a company name must never become a formula
        ws.row_dimensions[1].height = 30
        ws["A2"] = f"{BRAND} {PRODUCT}"
        ws["A2"].font = Font(size=11, color=GREY)
        ws["A3"] = f"Attendance report: {_period(start, end, dept)}"
        ws["A3"].font = Font(bold=True)

        for i, (label, _) in enumerate(cols, 1):
            c = ws.cell(header_row, i, label)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center")

        widths = [len(label) for label, _ in cols]
        for r, rec in enumerate(data, header_row + 1):
            for i, (_, key) in enumerate(cols, 1):
                value = rec[key]
                cell = ws.cell(r, i, value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"  # never let a name become a formula
                if isinstance(value, float):
                    cell.number_format = "0.00"
                if (key == "status" and value in FLAG) or (key in ("late_min", "absent") and value):
                    cell.font = Font(color=RED, bold=(key == "status"))
                widths[i - 1] = max(widths[i - 1], len(str(value)))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 3, 10), 34)
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        if data:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{header_row + len(data)}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb = Workbook()
    fill_sheet(wb.active, "Summary", SUMMARY_COLS, _summary_rows(summary))
    fill_sheet(wb.create_sheet(), "Daily detail", DETAIL_COLS, rows)
    wb.properties.title = f"{company} - Attendance report" if company else "4nextech Time - Attendance report"
    wb.properties.creator = "4nextech Time"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
_fonts_cache = None


def _fonts():
    """Prefer a Unicode TrueType font so names with letters like Ɔ or ɛ print correctly."""
    global _fonts_cache
    if _fonts_cache:
        return _fonts_cache
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pairs = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ]
    for regular, bold in pairs:
        if os.path.exists(regular) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont("NT-Regular", regular))
                pdfmetrics.registerFont(TTFont("NT-Bold", bold))
                _fonts_cache = ("NT-Regular", "NT-Bold")
                return _fonts_cache
            except Exception:  # noqa: BLE001
                continue
    _fonts_cache = ("Helvetica", "Helvetica-Bold")
    return _fonts_cache


def pdf_bytes(rows, summary, start, end, dept=None, company=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    regular, bold = _fonts()
    blue, grey, red = colors.HexColor("#" + BLUE), colors.HexColor("#" + GREY), colors.HexColor("#" + RED)
    page_w = landscape(A4)[0]
    margin = 12 * mm

    brand_style = ParagraphStyle("brand", fontName=bold, fontSize=26, leading=28, textColor=blue)
    product_style = ParagraphStyle("product", fontName=regular, fontSize=11, leading=13, textColor=grey)
    title_style = ParagraphStyle("title", fontName=bold, fontSize=13, leading=16, spaceBefore=10)
    sub_style = ParagraphStyle("sub", fontName=regular, fontSize=9, leading=12, textColor=grey, spaceAfter=6)
    head_style = ParagraphStyle("head", fontName=bold, fontSize=8, leading=10, textColor=colors.white)
    cell_style = ParagraphStyle("cell", fontName=regular, fontSize=8, leading=10)
    flag_style = ParagraphStyle("flag", parent=cell_style, fontName=bold, textColor=red)

    def table(cols, data, widths_mm, flag_keys=()):
        body = [[Paragraph(escape(label), head_style) for label, _ in cols]]
        for rec in data:
            line = []
            for _, key in cols:
                value = rec[key]
                is_flag = (key == "status" and value in FLAG) or (key in flag_keys and value)
                line.append(Paragraph(escape(str(value)), flag_style if is_flag else cell_style))
            body.append(line)
        t = Table(body, colWidths=[w * mm for w in widths_mm], repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), blue),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F7")]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D5DD")),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    def footer(canvas, doc_):
        canvas.saveState()
        y = 7 * mm
        canvas.setFont(bold, 8)
        canvas.setFillColor(blue)
        canvas.drawString(margin, y, BRAND)
        offset = canvas.stringWidth(BRAND, bold, 8) + 3
        canvas.setFont(regular, 8)
        canvas.setFillColor(grey)
        canvas.drawString(margin + offset, y, PRODUCT)
        canvas.drawCentredString(page_w / 2, y, f"Generated {datetime.now():%d %b %Y %H:%M}")
        canvas.drawRightString(page_w - margin, y, f"Page {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin,
        topMargin=11 * mm, bottomMargin=14 * mm,
        title=f"{company} - Attendance report" if company else "4nextech Time - Attendance report", author="4nextech Time",
    )
    story = [
        Paragraph(escape(company or BRAND), brand_style),
        Paragraph(f"{BRAND} {PRODUCT}", product_style),
        Paragraph("Attendance report", title_style),
        Paragraph(escape(_period(start, end, dept)), sub_style),
        Paragraph("Summary", title_style),
    ]
    summ = _summary_rows(summary)
    story.append(table(SUMMARY_COLS, summ, [16, 52, 32, 22, 18, 20, 32, 26, 28], flag_keys=("absent",))
                 if summ else Paragraph("No data for this period.", cell_style))
    story += [Spacer(1, 6), Paragraph("Daily detail", title_style)]
    story.append(table(DETAIL_COLS, rows, [22, 14, 42, 28, 22, 15, 15, 18, 14, 18, 20, 24], flag_keys=("late_min",))
                 if rows else Paragraph("No rows for this period.", cell_style))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
